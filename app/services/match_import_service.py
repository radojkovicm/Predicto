"""Bulk match import from an admin-filled Excel spreadsheet.

Workflow: admin downloads a template scoped to one competition (with its real
phase names baked in), fills it in by hand or via ChatGPT, uploads it back.
Only adds new fixtures — never touches existing matches or results. A stopgap
before a proper fixtures API integration.
"""
from datetime import datetime, timezone
from io import BytesIO
from typing import Optional

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from app.models.models import Match, Phase

TEMPLATE_HEADERS = [
    "Phase", "Kickoff UTC (YYYY-MM-DD HH:MM)", "Team 1 Name", "Team 1 Code",
    "Team 2 Name", "Team 2 Code",
]
EXAMPLE_ROW = [
    "(EXAMPLE — delete this row)", datetime(2026, 9, 16, 21, 0),
    "Real Madrid", "rma", "Liverpool", "liv",
]


def build_template_workbook(phases: list[Phase]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matches"
    ws.append(TEMPLATE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws.append(EXAMPLE_ROW)
    for cell in ws[2]:
        cell.font = Font(italic=True, color="999999")
    ws["B2"].number_format = "YYYY-MM-DD HH:MM"

    widths = [26, 26, 20, 12, 20, 12]
    for col, width in zip("ABCDEF", widths):
        ws.column_dimensions[col].width = width

    phases_sheet = wb.create_sheet("Valid Phases")
    phases_sheet.append(["Phase name (must match exactly)"])
    phases_sheet["A1"].font = Font(bold=True)
    phases_sheet.column_dimensions["A"].width = 30
    for phase in phases:
        phases_sheet.append([phase.name])

    if phases:
        dv = DataValidation(
            type="list",
            formula1=f"'Valid Phases'!$A$2:$A${len(phases) + 1}",
            allow_blank=True,
        )
        dv.error = "Pick a phase name from the 'Valid Phases' sheet (or type it exactly)."
        dv.errorTitle = "Unknown phase"
        ws.add_data_validation(dv)
        dv.add(f"A2:A1000")

    return wb


def _parse_kickoff_cell(value) -> Optional[datetime]:
    if isinstance(value, datetime):
        dt = value
    elif isinstance(value, str) and value.strip():
        dt = None
        for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M"):
            try:
                dt = datetime.strptime(value.strip(), fmt)
                break
            except ValueError:
                continue
        if dt is None:
            return None
    else:
        return None
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def parse_upload(db: Session, file_bytes: bytes, competition_id: int) -> tuple[int, list[str]]:
    """Adds new matches from the uploaded workbook to the given competition.
    Never touches existing matches. Returns (added_count, row-level messages).
    """
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb["Matches"] if "Matches" in wb.sheetnames else wb.active

    phases = db.query(Phase).filter(Phase.competition_id == competition_id).all()
    phases_by_name = {p.name.strip().lower(): p for p in phases}

    existing = db.query(Match).filter(Match.competition_id == competition_id).all()
    seen_keys = {
        (m.phase_id, m.team1_code, m.team2_code,
         m.kickoff_utc if m.kickoff_utc.tzinfo else m.kickoff_utc.replace(tzinfo=timezone.utc))
        for m in existing
    }

    added = 0
    messages: list[str] = []
    to_add = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue
        phase_name, kickoff_raw, t1_name, t1_code, t2_name, t2_code = (list(row) + [None] * 6)[:6]

        if not (phase_name and t1_name and t1_code and t2_name and t2_code and kickoff_raw):
            messages.append(f"Row {row_idx}: missing a required field — skipped.")
            continue

        phase = phases_by_name.get(str(phase_name).strip().lower())
        if not phase:
            messages.append(f"Row {row_idx}: phase '{phase_name}' not found in this competition — skipped.")
            continue

        kickoff = _parse_kickoff_cell(kickoff_raw)
        if not kickoff:
            messages.append(f"Row {row_idx}: couldn't read kickoff time '{kickoff_raw}' — skipped.")
            continue

        team1_code = str(t1_code).strip().lower()[:10]
        team2_code = str(t2_code).strip().lower()[:10]
        key = (phase.id, team1_code, team2_code, kickoff)
        if key in seen_keys:
            messages.append(f"Row {row_idx}: '{t1_name}' vs '{t2_name}' at that kickoff already exists — skipped.")
            continue
        seen_keys.add(key)

        to_add.append(Match(
            phase_id=phase.id,
            competition_id=competition_id,
            team1_code=team1_code,
            team1_name=str(t1_name).strip()[:50],
            team2_code=team2_code,
            team2_name=str(t2_name).strip()[:50],
            kickoff_utc=kickoff,
        ))
        added += 1

    db.add_all(to_add)
    db.commit()
    return added, messages
